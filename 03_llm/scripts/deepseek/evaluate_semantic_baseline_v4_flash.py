from __future__ import annotations

import argparse
import json
from collections import Counter
from pathlib import Path
from typing import Iterable

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


SCRIPT_DIR = Path(__file__).resolve().parent
ROOT_DIR = SCRIPT_DIR.parents[2]
DEEPSEEK_DIR = ROOT_DIR / "03_llm" / "deepseek"

INPUT_PATH = DEEPSEEK_DIR / "results" / "deepseek_v4_flash_semantic_baseline_1027.json"
SUMMARY_CSV_PATH = DEEPSEEK_DIR / "results" / "deepseek_v4_flash_semantic_baseline_summary.csv"
EVAL_XLSX_PATH = DEEPSEEK_DIR / "results" / "deepseek_v4_flash_semantic_baseline_evaluation.xlsx"
REPORT_PATH = DEEPSEEK_DIR / "results" / "deepseek_v4_flash_semantic_baseline_report.md"


def load_rows(path: Path) -> list[dict]:
    rows = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(rows, list):
        raise RuntimeError(f"{path} root must be a JSON list")
    return rows


def success_rows(rows: Iterable[dict]) -> list[dict]:
    return [row for row in rows if row.get("status") == "success" and row.get("final_label") in {"TP", "FP"}]


def failed_rows(rows: Iterable[dict]) -> list[dict]:
    return [row for row in rows if row.get("status") != "success" or row.get("final_label") not in {"TP", "FP"}]


def safe_ratio(num: int, den: int) -> float:
    return num / den if den else 0.0


def confusion_metrics(rows: Iterable[dict]) -> dict:
    rows = list(rows)
    tp = sum(1 for row in rows if row.get("ground_truth_label") == "TP" and row.get("final_label") == "TP")
    fp = sum(1 for row in rows if row.get("ground_truth_label") == "FP" and row.get("final_label") == "TP")
    tn = sum(1 for row in rows if row.get("ground_truth_label") == "FP" and row.get("final_label") == "FP")
    fn = sum(1 for row in rows if row.get("ground_truth_label") == "TP" and row.get("final_label") == "FP")
    total = tp + fp + tn + fn
    precision = safe_ratio(tp, tp + fp)
    recall = safe_ratio(tp, tp + fn)
    return {
        "Count": total,
        "TP": tp,
        "FP": fp,
        "TN": tn,
        "FN": fn,
        "Precision": precision,
        "Recall": recall,
        "F1-score": safe_ratio(2 * precision * recall, precision + recall),
        "Accuracy": safe_ratio(tp + tn, total),
    }


def expand_cwes(rows: list[dict]) -> pd.DataFrame:
    expanded: list[dict] = []
    for row in rows:
        cwes = row.get("cwe")
        if isinstance(cwes, list) and cwes:
            for cwe in cwes:
                expanded.append({"group": cwe, **row})
        else:
            expanded.append({"group": "UNAVAILABLE", **row})
    return pd.DataFrame(expanded)


def grouped_metrics(rows: list[dict], field_name: str) -> pd.DataFrame:
    if field_name == "cwe":
        group_df = expand_cwes(rows)
        group_key = "group"
    else:
        group_df = pd.DataFrame(rows)
        group_key = field_name
    output_rows: list[dict] = []
    if group_df.empty:
        return pd.DataFrame()
    for group_value, subset in group_df.groupby(group_key, dropna=False):
        metrics = confusion_metrics(subset.to_dict("records"))
        metrics[field_name] = group_value
        output_rows.append(metrics)
    order = [field_name, "Count", "TP", "FP", "TN", "FN", "Precision", "Recall", "F1-score", "Accuracy"]
    return pd.DataFrame(output_rows)[order].sort_values(by=["Count", field_name], ascending=[False, True]).reset_index(drop=True)


def prediction_distribution(rows: list[dict]) -> pd.DataFrame:
    gt_counter = Counter(row.get("ground_truth_label") for row in rows)
    pred_counter = Counter(row.get("final_label") for row in rows)
    return pd.DataFrame(
        [
            {"Metric": "Predicted TP", "Value": pred_counter.get("TP", 0)},
            {"Metric": "Predicted FP", "Value": pred_counter.get("FP", 0)},
            {"Metric": "Ground-truth TP", "Value": gt_counter.get("TP", 0)},
            {"Metric": "Ground-truth FP", "Value": gt_counter.get("FP", 0)},
        ]
    )


def run_metadata(rows: list[dict], failed: list[dict], input_path: Path) -> pd.DataFrame:
    request_models = sorted({row.get("request_model", "") for row in rows if row.get("request_model")})
    response_models = sorted({row.get("response_model", "") for row in rows if row.get("response_model")})
    return pd.DataFrame(
        [
            {"Field": "input_path", "Value": str(input_path)},
            {"Field": "row_count", "Value": len(rows)},
            {"Field": "success_count", "Value": len(success_rows(rows))},
            {"Field": "failed_count", "Value": len(failed)},
            {"Field": "request_models", "Value": ", ".join(request_models)},
            {"Field": "response_models", "Value": ", ".join(response_models)},
            {"Field": "thinking_mode", "Value": ", ".join(sorted({str(row.get('thinking_mode', '')) for row in rows if row.get('thinking_mode')}))},
            {"Field": "temperature", "Value": ", ".join(sorted({str(row.get('temperature', '')) for row in rows}))},
            {"Field": "top_p", "Value": ", ".join(sorted({str(row.get('top_p', '')) for row in rows}))},
            {"Field": "prompt_strategy", "Value": ", ".join(sorted({str(row.get('prompt_strategy', '')) for row in rows if row.get('prompt_strategy')}))},
            {"Field": "retry_total", "Value": sum(int(row.get("retry_count", 0) or 0) for row in rows)},
            {"Field": "repair_total", "Value": sum(int(row.get("repair_count", 0) or 0) for row in rows)},
        ]
    )


def save_excel(path: Path, sheets: dict[str, pd.DataFrame]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for sheet_name, frame in sheets.items():
            output = frame.copy()
            for column in output.columns:
                if output[column].dtype.kind == "f":
                    output[column] = output[column].round(6)
            output.to_excel(writer, sheet_name=sheet_name[:31], index=False)
    workbook = load_workbook(path)
    for worksheet in workbook.worksheets:
        for column in worksheet.columns:
            width = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column)
            worksheet.column_dimensions[get_column_letter(column[0].column)].width = min(width + 2, 80)
    workbook.save(path)


def build_report(
    *,
    input_path: Path,
    metrics: dict,
    rows: list[dict],
    failed: list[dict],
    tool_metrics: pd.DataFrame,
) -> str:
    gt_counter = Counter(row.get("ground_truth_label") for row in rows)
    pred_counter = Counter(row.get("final_label") for row in rows if row.get("status") == "success")
    lines = [
        "# DeepSeek V4-Flash Semantic Baseline Report",
        "",
        "## Run Summary",
        f"- Input: `{input_path}`",
        f"- Total rows: {len(rows)}",
        f"- Successful rows: {len(success_rows(rows))}",
        f"- Failed rows: {len(failed)}",
        f"- Retry total: {sum(int(row.get('retry_count', 0) or 0) for row in rows)}",
        f"- JSON repair total: {sum(int(row.get('repair_count', 0) or 0) for row in rows)}",
        "",
        "## Overall Metrics",
        f"- TP: {metrics['TP']}",
        f"- FP: {metrics['FP']}",
        f"- TN: {metrics['TN']}",
        f"- FN: {metrics['FN']}",
        f"- Precision: {metrics['Precision']:.6f}",
        f"- Recall: {metrics['Recall']:.6f}",
        f"- F1-score: {metrics['F1-score']:.6f}",
        f"- Accuracy: {metrics['Accuracy']:.6f}",
        "",
        "## Prediction Distribution",
        f"- Predicted TP: {pred_counter.get('TP', 0)}",
        f"- Predicted FP: {pred_counter.get('FP', 0)}",
        f"- Ground-truth TP: {gt_counter.get('TP', 0)}",
        f"- Ground-truth FP: {gt_counter.get('FP', 0)}",
        "",
        "## SCA-Level Snapshot",
    ]
    if tool_metrics.empty:
        lines.append("- No successful rows available.")
    else:
        for _, row in tool_metrics.iterrows():
            lines.append(
                f"- {row['tool_name']}: Count={int(row['Count'])}, TP={int(row['TP'])}, FP={int(row['FP'])}, TN={int(row['TN'])}, FN={int(row['FN'])}, F1={row['F1-score']:.6f}, Accuracy={row['Accuracy']:.6f}"
            )
    return "\n".join(lines) + "\n"


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-path", default=str(INPUT_PATH))
    parser.add_argument("--summary-csv-path", default=str(SUMMARY_CSV_PATH))
    parser.add_argument("--evaluation-xlsx-path", default=str(EVAL_XLSX_PATH))
    parser.add_argument("--report-path", default=str(REPORT_PATH))
    args = parser.parse_args()

    input_path = Path(args.input_path).resolve()
    summary_csv_path = Path(args.summary_csv_path).resolve()
    evaluation_xlsx_path = Path(args.evaluation_xlsx_path).resolve()
    report_path = Path(args.report_path).resolve()
    summary_csv_path.parent.mkdir(parents=True, exist_ok=True)
    evaluation_xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.parent.mkdir(parents=True, exist_ok=True)

    rows = load_rows(input_path)
    successful = success_rows(rows)
    failed = failed_rows(rows)
    if len(rows) != 1027:
        raise RuntimeError(f"Expected 1027 rows in output, found {len(rows)}")
    if len(successful) + len(failed) != len(rows):
        raise RuntimeError("Row accounting mismatch during evaluation")

    overall = confusion_metrics(successful)
    tool_level = grouped_metrics(successful, "tool_name")
    rule_level = grouped_metrics(successful, "rule_id")
    cwe_level = grouped_metrics(successful, "cwe")
    prediction_df = prediction_distribution(rows)
    metadata_df = run_metadata(rows, failed, input_path)
    failed_df = pd.DataFrame(
        [
            {
                "warning_id": row.get("warning_id"),
                "tool_name": row.get("tool_name"),
                "project_name": row.get("project_name"),
                "rule_id": row.get("rule_id"),
                "ground_truth_label": row.get("ground_truth_label"),
                "status": row.get("status"),
                "error_message": row.get("error_message"),
                "retry_count": row.get("retry_count"),
                "repair_count": row.get("repair_count"),
                "response_model": row.get("response_model"),
            }
            for row in failed
        ]
    )
    summary_rows = [
        {"Metric": key, "Value": value}
        for key, value in overall.items()
    ]
    pd.DataFrame(summary_rows).to_csv(summary_csv_path, index=False, encoding="utf-8-sig")

    save_excel(
        evaluation_xlsx_path,
        {
            "Overall": pd.DataFrame([overall]),
            "Confusion Matrix": pd.DataFrame(
                [
                    {"Ground Truth": "TP", "Pred TP": overall["TP"], "Pred FP": overall["FN"]},
                    {"Ground Truth": "FP", "Pred TP": overall["FP"], "Pred FP": overall["TN"]},
                ]
            ),
            "SCA Level": tool_level,
            "Rule Level": rule_level,
            "CWE Level": cwe_level,
            "Prediction Distribution": prediction_df,
            "Run Metadata": metadata_df,
            "Failed Samples": failed_df if not failed_df.empty else pd.DataFrame(columns=["warning_id"]),
        },
    )

    report_text = build_report(
        input_path=input_path,
        metrics=overall,
        rows=rows,
        failed=failed,
        tool_metrics=tool_level,
    )
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(report_text, encoding="utf-8")


if __name__ == "__main__":
    main()
